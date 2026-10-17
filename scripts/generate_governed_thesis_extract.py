"""Build the public, governed thesis evidence extract.

This replaces the earlier unreconciled public thesis draft.  Every quantitative
statement is derived from a named repository source and carries an evidence
status.  The university thesis itself remains withheld pending academic and
source reconciliation.
"""

from __future__ import annotations

import csv
import shutil
import tempfile
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

import matplotlib.pyplot as plt
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from scipy import stats


ROOT = Path(__file__).resolve().parents[1]
RAW_PATH = ROOT / "data" / "raw_movements.csv"
SITE_HOUR_PATH = ROOT / "data" / "site_hour.csv"
INCIDENT_PATH = ROOT / "data" / "incidents.csv"
OUTPUTS = (
    ROOT / "downloads" / "Final_Thesis.docx",
    ROOT / "website" / "public" / "downloads" / "Final_Thesis.docx",
    ROOT / "final_deliverables" / "Final_Thesis.docx",
)
BUILD_DIR = Path(tempfile.gettempdir()) / "tricycles-governed-thesis-build"
PRIMARY = RGBColor(31, 78, 121)
LIGHT_BLUE = "D9EAF7"
BLACK = RGBColor(0, 0, 0)
TABLE_COUNTER = 0
FIGURE_COUNTER = 0


NUMERIC_FIELDS = [
    "MotorizedTotal",
    "PassengerTricycle",
    "CargoTricycle",
    "TricycleTotal",
]


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def number(value: str) -> float:
    return float(value or 0)


def fmt(value: float, digits: int = 0) -> str:
    if digits == 0:
        return f"{value:,.0f}"
    return f"{value:,.{digits}f}"


def set_cell_fill(cell, colour: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), colour)


def set_cell_margins(cell, top=80, start=80, bottom=80, end=80) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Arial", size=9, bold=False, colour=BLACK) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = colour


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for name, size in (("Title", 18), ("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 10.5)):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)

    header = section.header.paragraphs[0]
    header.text = "TRICYCLE PCU RESEARCH: CONTROLLED EVIDENCE EXTRACT"
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in header.runs:
        set_run_font(run, size=7, bold=True, colour=PRIMARY)
    add_page_number(section.footer.paragraphs[0])


def add_title_page(doc: Document) -> None:
    for line in (
        "KAMPALA INTERNATIONAL UNIVERSITY",
        "DIRECTORATE OF HIGHER DEGREES AND RESEARCH",
        "SCHOOL OF ENGINEERING AND APPLIED SCIENCES",
        "DEPARTMENT OF CIVIL ENGINEERING",
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        set_run_font(run, size=11, bold=True)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("ANALYSIS OF TRICYCLE PASSENGER CAR UNIT VALUES FOR ENHANCED TRAFFIC FLOW IN KAMPALA CITY, UGANDA")
    set_run_font(run, size=18, bold=True, colour=PRIMARY)
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(18)

    for line, size, bold in (
        ("CONTROLLED GOVERNED EVIDENCE EXTRACT", 14, True),
        ("Public replacement for the unreconciled thesis draft", 10, False),
        ("SSERUNJOGI AMBROSE", 11, True),
        ("Student number: 2022-08-12043", 10, False),
        ("Evidence cut-off: 18 October 2026", 10, False),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        set_run_font(run, size=size, bold=bold)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        "STATUS: This document is an evidence extract, not the final university thesis. "
        "The earlier full thesis remains withheld until every claim, table, test and figure is reconciled to a governed source."
    )
    set_run_font(run, size=10, bold=True, colour=RGBColor(156, 87, 0))
    p.paragraph_format.space_before = Pt(20)
    doc.add_page_break()


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_body(doc: Document, text: str, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_run_font(first, size=10, bold=True)
        second = p.add_run(text[len(bold_prefix):])
        set_run_font(second, size=10)
    else:
        run = p.add_run(text)
        set_run_font(run, size=10)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], title: str, widths: list[float] | None = None) -> None:
    global TABLE_COUNTER
    TABLE_COUNTER += 1
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_fill(cell, LIGHT_BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(header)
        set_run_font(run, size=8, bold=True)
        if widths:
            cell.width = Inches(widths[index])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        if row_index % 2:
            for cell in cells:
                set_cell_fill(cell, "F7F7F7")
        for col_index, value in enumerate(values):
            cell = cells[col_index]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            run = p.add_run(str(value))
            set_run_font(run, size=7.5)
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_with_next = True
    run = caption.add_run(f"Table {TABLE_COUNTER}: {title}")
    set_run_font(run, size=8, bold=True)


def add_figure(doc: Document, path: Path, title: str) -> None:
    global FIGURE_COUNTER
    FIGURE_COUNTER += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(6.25))
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_with_next = True
    run = caption.add_run(f"Figure {FIGURE_COUNTER}: {title}")
    set_run_font(run, size=8, bold=True)


def aggregate(rows: list[dict[str, str]], field: str) -> list[list[str]]:
    grouped: dict[str, dict[str, float]] = defaultdict(lambda: {key: 0 for key in NUMERIC_FIELDS} | {"records": 0})
    for row in rows:
        category = str(row.get(field, "")).strip() or "Not recorded"
        grouped[category]["records"] += 1
        for key in NUMERIC_FIELDS:
            grouped[category][key] += number(row[key])
    result = []
    for category in sorted(grouped, key=lambda value: value.casefold()):
        item = grouped[category]
        share = item["TricycleTotal"] / item["MotorizedTotal"] if item["MotorizedTotal"] else 0
        result.append([
            category,
            fmt(item["records"]),
            fmt(item["MotorizedTotal"]),
            fmt(item["PassengerTricycle"]),
            fmt(item["CargoTricycle"]),
            fmt(item["TricycleTotal"]),
            f"{share * 100:.1f}%",
        ])
    return result


def trend_rows(rows: list[dict[str, str]], mode: str) -> list[list[str]]:
    copied = []
    for row in rows:
        item = dict(row)
        date = datetime.fromisoformat(row["Date"])
        if mode == "annual":
            item["Period"] = str(date.year)
        elif mode == "monthly":
            item["Period"] = date.strftime("%Y-%m")
        else:
            item["Period"] = f"FY{date.year}/{str(date.year + 1)[-2:]}" if date.month >= 7 else f"FY{date.year - 1}/{str(date.year)[-2:]}"
        copied.append(item)
    return aggregate(copied, "Period")


def weighted_mean(rows: list[dict[str, str]], value_field: str, weight_field: str) -> float:
    numerator = sum(number(row[value_field]) * number(row[weight_field]) for row in rows)
    denominator = sum(number(row[weight_field]) for row in rows)
    return numerator / denominator if denominator else 0


def make_charts(raw_rows, site_rows) -> tuple[Path, Path]:
    BUILD_DIR.mkdir(parents=True, exist_ok=True)
    site_data = aggregate(raw_rows, "Site")
    sites = [row[0] for row in site_data]
    totals = [float(row[5].replace(",", "")) for row in site_data]
    plt.figure(figsize=(8.0, 4.2))
    bars = plt.bar(sites, totals, color="#1F4E79")
    plt.ylabel("Measured tricycle count")
    plt.xlabel("Study site")
    plt.title("Complete measured tricycle population by site")
    plt.grid(axis="y", alpha=0.2)
    plt.bar_label(bars, labels=[f"{value:,.0f}" for value in totals], padding=3, fontsize=8)
    plt.tight_layout()
    measured = BUILD_DIR / "measured_tricycle_population.png"
    plt.savefig(measured, dpi=180)
    plt.close()

    grouped = defaultdict(list)
    for row in site_rows:
        grouped[row["Site"]].append(row)
    sites = sorted(grouped)
    passenger = [weighted_mean(grouped[site], "PassengerDynamicPCU", "PassengerTricycle") for site in sites]
    cargo = [weighted_mean(grouped[site], "CargoDynamicPCU", "CargoTricycle") for site in sites]
    x = list(range(len(sites)))
    width = 0.36
    plt.figure(figsize=(8.0, 4.2))
    plt.bar([value - width / 2 for value in x], passenger, width, label="Passenger tricycle", color="#2E75B6")
    plt.bar([value + width / 2 for value in x], cargo, width, label="Cargo tricycle", color="#70AD47")
    plt.xticks(x, sites)
    plt.ylabel("Modelled dynamic PCU")
    plt.xlabel("Study site")
    plt.title("Modelled dynamic PCU by site")
    plt.legend(frameon=False)
    plt.grid(axis="y", alpha=0.2)
    plt.tight_layout()
    modelled = BUILD_DIR / "modelled_dynamic_pcu.png"
    plt.savefig(modelled, dpi=180)
    plt.close()
    return measured, modelled


def build() -> None:
    raw_rows = read_csv(RAW_PATH)
    site_rows = read_csv(SITE_HOUR_PATH)
    incident_rows = read_csv(INCIDENT_PATH)
    measured_chart, modelled_chart = make_charts(raw_rows, site_rows)

    doc = Document()
    configure_document(doc)
    add_title_page(doc)

    add_heading(doc, "Document control", 1)
    add_body(doc, "This controlled extract removes unsupported claims from the public download while preserving the study topic and the validated evidence that can be reproduced from repository sources.")
    add_table(doc, ["Field", "Controlled value"], [
        ["Document status", "Governed evidence extract, not the final thesis"],
        ["Public filename", "Final_Thesis.docx, retained for link stability"],
        ["Prepared", datetime.now(timezone.utc).strftime("%d %B %Y, %H:%M UTC")],
        ["Language", "UK English"],
        ["Coverage", "Five Kampala study sites, 2 and 3 August 2026, 07:00 to 18:00"],
        ["Superseded material", "Earlier unreconciled thesis narrative, random appendices and unsupported simulation claims"],
    ], "Document Control")
    add_table(doc, ["Evidence status", "Meaning in this document"], [
        ["Measured", "Direct aggregation of the governed repository extract, with no Top-N filtering."],
        ["Modelled", "Calculated from measured site-hour rows using disclosed, editable assumptions."],
        ["Unavailable", "Required dimension or analysis not supported by the governed source and not fabricated."],
    ], "Evidence Status Definitions")

    add_heading(doc, "Executive summary", 1)
    total_motorised = sum(number(row["MotorizedTotal"]) for row in raw_rows)
    total_passenger = sum(number(row["PassengerTricycle"]) for row in raw_rows)
    total_cargo = sum(number(row["CargoTricycle"]) for row in raw_rows)
    add_body(doc, f"Measured evidence: the governed extract contains {fmt(len(raw_rows))} classified movement records from five study sites and {fmt(len(site_rows))} complete site-hour rows. Across this population, {fmt(total_motorised)} motorised vehicles were recorded, including {fmt(total_passenger)} passenger tricycles and {fmt(total_cargo)} cargo tricycles. These figures are descriptive totals for the supplied study period and are not network-wide estimates.")
    add_body(doc, "Modelled evidence: passenger and cargo PCU values are derived from a speed-area equivalency formula and a demand-pressure adjustment. Vehicle dimensions, speeds and the demand sensitivity coefficient are assumptions recorded in the governed analytical code, not measured field observations.")
    add_body(doc, "Unavailable evidence: the governed tricycle counts do not contain Surface Type, Functional Class, Road Type or Sub-Region. Wet-weather observations, measured headways, measured speeds, capacity, saturation flow, queue length, signal timing outcomes and VISSIM calibration records are also unavailable. This extract does not report those items as findings.")

    add_heading(doc, "1. Scope and governance", 1)
    add_heading(doc, "1.1 Purpose", 2)
    add_body(doc, "The purpose of this extract is to publish the reproducible evidence currently available for the tricycle PCU study without presenting unverified material as completed research. The full university thesis remains withheld until its literature, methodology, results, appendices and narrative claims have been independently reconciled.")
    add_heading(doc, "1.2 Named governed sources", 2)
    source_rows = [
        ["data/raw_movements.csv", "Measured", f"{fmt(len(raw_rows))} classified movement records; retains every included record and source-workbook name."],
        ["data/site_hour.csv", "Measured plus modelled fields", f"{fmt(len(site_rows))} complete site-hour rows; dynamic PCU fields are explicitly modelled."],
        ["data/incidents.csv", "Measured extract", f"{fmt(len(incident_rows))} incident type and severity records."],
        ["website/src/lib/trafficStats.js", "Method", "Governed PCU assumptions, formulas, category definitions and statistical methods."],
        ["docs/SSERUNJOGI AMBROSE_Proposal_V2.1.pdf", "Proposal reference", "Study rationale and literature context only; not treated as observed results."],
    ]
    add_table(doc, ["Repository source", "Evidence class", "Use in this extract"], source_rows, "Governed Source Register")
    source_files = sorted({row["SourceFile"] for row in raw_rows})
    add_body(doc, "The movement extract records five source-workbook names: " + "; ".join(source_files) + ". The public repository does not contain those original workbooks under these names, so the CSV extract is the controlled public source used here. This limitation prevents a fresh workbook-to-record reconciliation within the public package.")

    add_heading(doc, "1.3 Required dimension coverage", 2)
    add_table(doc, ["Required dimension", "Status", "Governance treatment"], [
        ["Surface Type", "Unavailable", "Not present; no Paved or Unpaved allocation is fabricated."],
        ["Functional Class", "Unavailable", "Not present; not fabricated."],
        ["Road Type", "Unavailable", "Not present; not fabricated."],
        ["Sub-Region", "Unavailable", "Not present; not fabricated."],
        ["Financial Year", "Available", "Derived from the recorded dates using July to June, giving FY2026/27."],
        ["Annual trend", "Available, one period", "All observed records fall in 2026; no other year is inferred."],
        ["Monthly trend", "Available, one period", "All observed records fall in August 2026; no other month is inferred."],
        ["Other categorical fields", "Available", "Every category in the governed movement extract is reported in Appendix A."],
    ], "Required Reporting Dimension Coverage")

    add_heading(doc, "2. Method", 1)
    add_heading(doc, "2.1 Measured population", 2)
    add_body(doc, "All rows with IncludedInAnalysis equal to Yes are retained. No Top-N filter is applied. Seven high-concentration review flags remain in the primary population exactly as recorded. Category labels are preserved, with blank labels displayed as Not recorded.")
    add_table(doc, ["Measure", "Value", "Status"], [
        ["Movement records", fmt(len(raw_rows)), "Measured"],
        ["Site-hour rows", fmt(len(site_rows)), "Measured aggregation"],
        ["Study sites", fmt(len({row['Site'] for row in raw_rows})), "Measured"],
        ["Incident records", fmt(len(incident_rows)), "Measured extract"],
        ["Observation dates", fmt(len({row['Date'] for row in raw_rows})), "Measured"],
        ["Observation hours", "07:00 to 18:00", "Measured window"],
    ], "Governed Study Population")

    add_heading(doc, "2.2 PCU model", 2)
    add_body(doc, "Static PCU is modelled as (car speed / subject vehicle speed) x (subject vehicle area / car area), where area equals length x width. The governed assumptions are: car 4.0 m x 1.65 m at 45 km/h; passenger tricycle 2.9 m x 1.4 m at 27.5 km/h; cargo tricycle 3.2 m x 1.5 m at 27.5 km/h.")
    add_body(doc, "Dynamic PCU is modelled as Static PCU x [1 + beta x (Demand Pressure - 0.5)], where beta equals 0.4 and Demand Pressure equals hourly motorised flow divided by the maximum recorded hourly flow at the same site. This relationship is definitional and is not reported as an independently observed causal effect.")
    add_table(doc, ["Parameter", "Value", "Status"], [
        ["Passenger tricycle static PCU", "1.0066", "Modelled assumption"],
        ["Cargo tricycle static PCU", "1.1901", "Modelled assumption"],
        ["Demand sensitivity beta", "0.4", "Modelled assumption"],
        ["Wet-pavement multiplier", "1.10", "Sensitivity assumption only, not used as a measured finding"],
    ], "PCU Model Assumptions")

    add_heading(doc, "2.3 Statistical method", 2)
    add_body(doc, "One-way ANOVA and Kruskal-Wallis tests compare tricycle counts across the complete site and time-period populations. Pearson correlation evaluates paired hourly motorised flow and tricycle totals. The incident chi-square test evaluates the complete Incident Type by Severity table. These tests describe association or distributional difference and do not establish causality.")

    add_heading(doc, "3. Results", 1)
    add_heading(doc, "3.1 Complete measured population by site", 2)
    common_headers = ["Category", "Records", "Motorised vehicles", "Passenger tricycles", "Cargo tricycles", "Total tricycles", "Tricycle share"]
    add_table(doc, common_headers, aggregate(raw_rows, "Site"), "Complete Measured Traffic Population by Study Site")
    add_figure(doc, measured_chart, "Complete Measured Tricycle Population by Study Site")

    grouped_sites = defaultdict(list)
    for row in site_rows:
        grouped_sites[row["Site"]].append(row)
    modelled_site_rows = []
    for site in sorted(grouped_sites):
        rows = grouped_sites[site]
        modelled_site_rows.append([
            site,
            rows[0]["JunctionType"],
            fmt(sum(number(row["TricycleTotal"]) for row in rows)),
            fmt(sum(number(row["DemandPressure"]) for row in rows) / len(rows), 3),
            fmt(max(number(row["DemandPressure"]) for row in rows), 3),
            fmt(weighted_mean(rows, "PassengerDynamicPCU", "PassengerTricycle"), 3),
            fmt(weighted_mean(rows, "CargoDynamicPCU", "CargoTricycle"), 3),
            fmt(weighted_mean(rows, "CombinedDynamicPCU", "TricycleTotal"), 3),
        ])
    add_heading(doc, "3.2 Modelled PCU by site", 2)
    add_table(doc, ["Site", "Junction type", "Measured tricycles", "Mean demand pressure", "Peak demand pressure", "Passenger PCU, modelled", "Cargo PCU, modelled", "Combined PCU, modelled"], modelled_site_rows, "Modelled Dynamic PCU by Study Site")
    add_figure(doc, modelled_chart, "Modelled Dynamic PCU by Study Site")

    grouped_periods = defaultdict(list)
    for row in site_rows:
        grouped_periods[row["TimePeriod"]].append(row)
    modelled_period_rows = []
    for period in sorted(grouped_periods):
        rows = grouped_periods[period]
        modelled_period_rows.append([
            period,
            fmt(len(rows)),
            fmt(sum(number(row["TricycleTotal"]) for row in rows)),
            fmt(sum(number(row["DemandPressure"]) for row in rows) / len(rows), 3),
            fmt(weighted_mean(rows, "PassengerDynamicPCU", "PassengerTricycle"), 3),
            fmt(weighted_mean(rows, "CargoDynamicPCU", "CargoTricycle"), 3),
            fmt(weighted_mean(rows, "CombinedDynamicPCU", "TricycleTotal"), 3),
        ])
    add_heading(doc, "3.3 Modelled PCU by time period", 2)
    add_table(doc, ["Time period", "Site-hour rows", "Measured tricycles", "Mean demand pressure", "Passenger PCU, modelled", "Cargo PCU, modelled", "Combined PCU, modelled"], modelled_period_rows, "Modelled Dynamic PCU by Time Period")

    add_heading(doc, "3.4 Reproducible statistical results", 2)
    site_groups = [[number(row["TricycleTotal"]) for row in raw_rows if row["Site"] == site] for site in sorted({row["Site"] for row in raw_rows})]
    period_groups = [[number(row["TricycleTotal"]) for row in raw_rows if row["TimePeriod"] == period] for period in sorted({row["TimePeriod"] for row in raw_rows})]
    site_anova = stats.f_oneway(*site_groups)
    site_kw = stats.kruskal(*site_groups)
    period_anova = stats.f_oneway(*period_groups)
    period_kw = stats.kruskal(*period_groups)
    correlation = stats.pearsonr([number(row["MotorizedFlow"]) for row in site_rows], [number(row["TricycleTotal"]) for row in site_rows])
    types = sorted({row["IncidentType"] for row in incident_rows})
    severities = sorted({row["Severity"] for row in incident_rows})
    incident_matrix = [[sum(1 for row in incident_rows if row["IncidentType"] == kind and row["Severity"] == severity) for severity in severities] for kind in types]
    chi2, incident_p, incident_df, _ = stats.chi2_contingency(incident_matrix)
    add_table(doc, ["Test", "Complete population", "Statistic", "p-value", "Evidence status", "Bounded interpretation"], [
        ["One-way ANOVA by site", f"n={len(raw_rows)}, 5 sites", f"F={site_anova.statistic:.3f}", f"{site_anova.pvalue:.3g}", "Measured counts", "Tests mean differences; no causal claim."],
        ["Kruskal-Wallis by site", f"n={len(raw_rows)}, 5 sites", f"H={site_kw.statistic:.3f}", f"{site_kw.pvalue:.3g}", "Measured counts", "Rank-based robustness check."],
        ["One-way ANOVA by period", f"n={len(raw_rows)}, 3 periods", f"F={period_anova.statistic:.3f}", f"{period_anova.pvalue:.3g}", "Measured counts", "Tests mean differences; no causal claim."],
        ["Kruskal-Wallis by period", f"n={len(raw_rows)}, 3 periods", f"H={period_kw.statistic:.3f}", f"{period_kw.pvalue:.3g}", "Measured counts", "Rank-based robustness check."],
        ["Pearson flow correlation", f"n={len(site_rows)} site-hours", f"r={correlation.statistic:.3f}", f"{correlation.pvalue:.3g}", "Measured paired totals", "Association only."],
        ["Incident type x severity", f"n={len(incident_rows)} incidents", f"chi-square={chi2:.3f}, df={incident_df}", f"{incident_p:.3g}", "Measured extract", "Association only; no exposure denominator."],
    ], "Reproducible Statistical Tests")

    add_heading(doc, "3.5 Complete annual, monthly and financial-year trends", 2)
    add_table(doc, common_headers, trend_rows(raw_rows, "annual"), "Complete Annual Traffic Trend")
    add_table(doc, common_headers, trend_rows(raw_rows, "monthly"), "Complete Monthly Traffic Trend")
    add_table(doc, common_headers, trend_rows(raw_rows, "financial_year"), "Complete Financial-Year Traffic Trend")
    add_body(doc, "All governed observations fall in August 2026 and FY2026/27. The one-row trend tables are the complete available population. No missing month or year is estimated or projected.")

    add_heading(doc, "4. Conclusions and evidence boundary", 1)
    add_body(doc, "The controlled evidence supports descriptive reporting of the supplied two-day count population, source-reconciled site-hour aggregation, a disclosed modelled PCU framework, and reproducible distributional tests. It does not support claims about a 20-day survey, measured speed or headway, capacity change, queue reduction, signal optimisation, wet-weather effects, VISSIM validation, economic benefits or network-wide traffic conditions.")
    add_body(doc, "The modelled PCU values are analytical scenarios rather than empirically calibrated PCU measurements. Before a final thesis is released, the original source workbooks, field protocols, enumerator records, speed and headway observations, calibration files, literature citations and academic approvals must be supplied and reconciled.")

    add_heading(doc, "Appendix A: Complete governed categorical breakdowns", 1)
    add_body(doc, "Every categorical field in data/raw_movements.csv is reported below. No category is suppressed. Date is reported in addition to the annual, monthly and financial-year tables above.")
    categorical_fields = [
        ("Site", "Study Site"),
        ("JunctionType", "Junction Type"),
        ("SourceFile", "Recorded Source Workbook"),
        ("Sheet", "Source Sheet"),
        ("Date", "Observation Date"),
        ("DateImputed", "Date Imputation Status"),
        ("DirFrom", "Origin Direction"),
        ("DirTo", "Destination Direction"),
        ("Weather", "Weather Label"),
        ("TimeFrom", "Observation Start Hour"),
        ("TimeTo", "Observation End Hour"),
        ("HighConcentrationReview", "High-Concentration Review Flag"),
        ("QCWarning", "Quality-Control Warning"),
        ("IncludedInAnalysis", "Analysis Inclusion Status"),
        ("TimePeriod", "Time Period"),
    ]
    for field, label in categorical_fields:
        add_heading(doc, label, 2)
        add_table(doc, common_headers, aggregate(raw_rows, field), f"Complete Measured Traffic Breakdown by {label}")

    add_heading(doc, "Appendix B: Complete incident breakdown", 1)
    incident_counts = defaultdict(lambda: {severity: 0 for severity in severities})
    for row in incident_rows:
        incident_counts[row["IncidentType"]][row["Severity"]] += 1
    incident_table = []
    for incident_type in types:
        counts = incident_counts[incident_type]
        incident_table.append([incident_type, *[fmt(counts[severity]) for severity in severities], fmt(sum(counts.values()))])
    add_table(doc, ["Incident Type", *severities, "Total"], incident_table, "Complete Incident Type by Severity Population")

    add_heading(doc, "Appendix C: Release conditions for the full thesis", 1)
    conditions = [
        "Original traffic-count workbooks must be present and reconciled record-for-record to data/raw_movements.csv.",
        "Measured speed and headway observations must be supplied before field-derived PCU claims can be made.",
        "Wet-weather, capacity, queue, signal timing and simulation claims require their own governed observations and methods.",
        "Every table and figure must cite its named source and state whether the values are measured, modelled, estimated, assigned, projected or unavailable.",
        "The literature and reference list must be verified against accessible primary sources.",
        "The final academic thesis must receive the required university review and approval before being labelled final.",
    ]
    for item in conditions:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run, size=10)

    for output in OUTPUTS:
        output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUTS[0])
    for output in OUTPUTS[1:]:
        shutil.copy2(OUTPUTS[0], output)
    print(f"Generated governed evidence extract: {OUTPUTS[0]}")
    print(f"Records: {len(raw_rows):,}; site-hours: {len(site_rows):,}; incidents: {len(incident_rows):,}")
    print(f"Tables: {TABLE_COUNTER}; figures: {FIGURE_COUNTER}")


if __name__ == "__main__":
    build()
